"""Generate sample resume documents (.txt, .docx, .pdf) for test uploading."""

import io
import sys
from pathlib import Path

# Add project root to sys.path
sys.path.insert(0, str(Path(__file__).resolve().parent.parent))

import docx
import pypdf

from app.samples import SAMPLE_PROFILES

SAMPLES_DIR = Path(__file__).resolve().parent / "sample_resumes"
SAMPLES_DIR.mkdir(parents=True, exist_ok=True)


def generate_documents():
    for profile in SAMPLE_PROFILES:
        base_name = profile["id"].replace("-", "_")
        text_content = profile["resume_text"].strip()

        # 1. Generate TXT
        txt_path = SAMPLES_DIR / f"{base_name}.txt"
        with open(txt_path, "w", encoding="utf-8") as f:
            f.write(text_content)
        print(f"Generated: {txt_path.name}")

        # 2. Generate DOCX
        docx_path = SAMPLES_DIR / f"{base_name}.docx"
        doc = docx.Document()
        lines = text_content.splitlines()
        if lines:
            doc.add_heading(lines[0], level=1)
            for line in lines[1:]:
                stripped = line.strip()
                if stripped.startswith("- "):
                    doc.add_paragraph(stripped[2:], style='List Bullet')
                elif stripped.isupper() and len(stripped) < 40:
                    doc.add_heading(stripped, level=2)
                elif stripped:
                    doc.add_paragraph(stripped)
        doc.save(str(docx_path))
        print(f"Generated: {docx_path.name}")

        # 3. Generate PDF
        try:
            from reportlab.lib.pagesizes import letter
            from reportlab.pdfgen import canvas

            pdf_path = SAMPLES_DIR / f"{base_name}.pdf"
            c = canvas.Canvas(str(pdf_path), pagesize=letter)
            width, height = letter
            y = height - 50

            c.setFont("Helvetica-Bold", 16)
            if lines:
                c.drawString(50, y, lines[0].strip())
                y -= 25

            c.setFont("Helvetica", 10)
            for line in lines[1:]:
                stripped = line.strip()
                if not stripped:
                    y -= 8
                    continue
                if stripped.isupper() and len(stripped) < 40:
                    y -= 10
                    c.setFont("Helvetica-Bold", 12)
                    c.drawString(50, y, stripped)
                    y -= 15
                    c.setFont("Helvetica", 10)
                else:
                    c.drawString(50, y, stripped[:95])
                    y -= 14

                if y < 50:
                    c.showPage()
                    c.setFont("Helvetica", 10)
                    y = height - 50

            c.save()
            print(f"Generated: {pdf_path.name}")
        except Exception as e:
            print(f"PDF generation error for {base_name}: {e}")

    print("All sample resume documents generated successfully!")


if __name__ == "__main__":
    generate_documents()
