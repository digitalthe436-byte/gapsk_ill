"""Resume Document Parser.

Extracts plain text reliably from multiple document variants (.pdf, .docx, .txt).
Handles file buffers, raw byte streams, and file paths.
"""

import io
from pathlib import Path
from typing import Union, BinaryIO

import pypdf
import docx

from backend.parsers.text_cleaner import clean_text


class DocumentParsingError(Exception):
    """Raised when parsing fails or the format is corrupted."""
    pass


def extract_text_from_pdf(file_source: Union[str, Path, BinaryIO, bytes]) -> str:
    """Extract plain text from a PDF file using pypdf.

    Args:
        file_source: File path, file-like object, or bytes.

    Returns:
        Extracted text string.
    """
    try:
        if isinstance(file_source, (str, Path)):
            reader = pypdf.PdfReader(str(file_source))
        elif isinstance(file_source, bytes):
            reader = pypdf.PdfReader(io.BytesIO(file_source))
        else:
            reader = pypdf.PdfReader(file_source)

        text_parts = []
        for idx, page in enumerate(reader.pages):
            page_text = page.extract_text()
            if page_text:
                text_parts.append(page_text)

        raw_text = "\n".join(text_parts)
        if not raw_text.strip():
            raise DocumentParsingError("PDF appears empty or contains only unscannable images.")
        return clean_text(raw_text)
    except Exception as e:
        if isinstance(e, DocumentParsingError):
            raise
        raise DocumentParsingError(f"Failed to extract text from PDF: {str(e)}") from e


def extract_text_from_docx(file_source: Union[str, Path, BinaryIO, bytes]) -> str:
    """Extract plain text from a Word DOCX file using python-docx.

    Args:
        file_source: File path, file-like object, or bytes.

    Returns:
        Extracted text string.
    """
    try:
        if isinstance(file_source, (str, Path)):
            doc = docx.Document(str(file_source))
        elif isinstance(file_source, bytes):
            doc = docx.Document(io.BytesIO(file_source))
        else:
            doc = docx.Document(file_source)

        text_parts = []
        # Paragraphs
        for paragraph in doc.paragraphs:
            if paragraph.text.strip():
                text_parts.append(paragraph.text)

        # Tables (often used in resumes for skills and experience columns)
        for table in doc.tables:
            for row in table.rows:
                row_cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if row_cells:
                    # Deduplicate repeated text from merged cells
                    unique_cells = []
                    for c in row_cells:
                        if not unique_cells or c != unique_cells[-1]:
                            unique_cells.append(c)
                    text_parts.append(" | ".join(unique_cells))

        raw_text = "\n".join(text_parts)
        if not raw_text.strip():
            raise DocumentParsingError("DOCX document appears empty.")
        return clean_text(raw_text)
    except Exception as e:
        if isinstance(e, DocumentParsingError):
            raise
        raise DocumentParsingError(f"Failed to extract text from DOCX: {str(e)}") from e


def extract_text_from_txt(file_source: Union[str, Path, BinaryIO, bytes]) -> str:
    """Extract and decode plain text."""
    try:
        if isinstance(file_source, (str, Path)):
            with open(file_source, "r", encoding="utf-8", errors="replace") as f:
                content = f.read()
        elif isinstance(file_source, bytes):
            content = file_source.decode("utf-8", errors="replace")
        else:
            raw = file_source.read()
            if isinstance(raw, bytes):
                content = raw.decode("utf-8", errors="replace")
            else:
                content = str(raw)
        return clean_text(content)
    except Exception as e:
        raise DocumentParsingError(f"Failed to read plain text file: {str(e)}") from e


def parse_resume(
    file_source: Union[str, Path, BinaryIO, bytes],
    filename: str = ""
) -> str:
    """Master resume parsing dispatcher based on file extension or filename.

    Args:
        file_source: File path, file-like object, or raw bytes.
        filename: Optional filename hint to determine extension.

    Returns:
        Cleaned extracted text string.
    """
    ext = ""
    if filename:
        ext = Path(filename).suffix.lower()
    elif isinstance(file_source, (str, Path)):
        ext = Path(file_source).suffix.lower()

    if ext == ".pdf":
        return extract_text_from_pdf(file_source)
    elif ext in {".docx", ".doc"}:
        return extract_text_from_docx(file_source)
    elif ext in {".txt", ""}:
        # Try PDF first if header matches %PDF
        if isinstance(file_source, bytes) and file_source.startswith(b"%PDF"):
            return extract_text_from_pdf(file_source)
        # Try DOCX if zip header matches PK
        if isinstance(file_source, bytes) and file_source.startswith(b"PK"):
            try:
                return extract_text_from_docx(file_source)
            except Exception:
                pass
        return extract_text_from_txt(file_source)
    else:
        raise DocumentParsingError(f"Unsupported file format '{ext}'. Allowed: .pdf, .docx, .txt")
