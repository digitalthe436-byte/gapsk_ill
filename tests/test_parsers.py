"""Unit tests for Data Ingestion & Text Extraction Parsers."""

import io
import unittest
from app.parsers.text_cleaner import clean_text, strip_html, normalize_unicode, segment_job_description
from app.parsers.resume_parser import parse_resume, extract_text_from_pdf, extract_text_from_docx, extract_text_from_txt

import pypdf
import docx


class TestTextCleaner(unittest.TestCase):

    def test_strip_html(self):
        raw = "<div><h1>Senior Software Engineer</h1><p>Must know <b>Python</b> &amp; Docker.</p></div>"
        cleaned = strip_html(raw)
        self.assertNotIn("<div>", cleaned)
        self.assertNotIn("<p>", cleaned)
        self.assertIn("Senior Software Engineer", cleaned)
        self.assertIn("Python", cleaned)

    def test_normalize_unicode(self):
        raw = "• Experienced with Python—and “smart quotes”\u00a0and tabs\t"
        normalized = normalize_unicode(raw)
        self.assertIn('"smart quotes"', normalized)
        self.assertIn("Python-and", normalized)

    def test_clean_text(self):
        raw = "   <div>Python   Developer\n\n\n\nWith    5   years exp.</div>  "
        cleaned = clean_text(raw)
        self.assertEqual(cleaned, "Python Developer\n\nWith 5 years exp.")

    def test_segment_job_description(self):
        jd = """
        About Us:
        Leading software tech company.

        Required Qualifications:
        - 3+ years experience with Python and PostgreSQL.
        - Solid experience with Docker.

        Preferred Qualifications:
        - Experience with Kubernetes.
        - Knowledge of AWS.
        """
        segmented = segment_job_description(jd)
        self.assertIn("Python", segmented["required"])
        self.assertIn("PostgreSQL", segmented["required"])
        self.assertIn("Kubernetes", segmented["preferred"])
        self.assertIn("AWS", segmented["preferred"])


class TestDocumentParsers(unittest.TestCase):

    def test_txt_parser(self):
        sample_txt = "John Doe\nSoftware Engineer with Python and React."
        extracted = extract_text_from_txt(sample_txt.encode("utf-8"))
        self.assertIn("John Doe", extracted)
        self.assertIn("Python and React", extracted)

    def test_docx_parser(self):
        # Create minimal in-memory DOCX
        doc = docx.Document()
        doc.add_heading("Jane Smith Resume", level=1)
        doc.add_paragraph("Full Stack Developer specializing in TypeScript and Node.js.")
        table = doc.add_table(rows=1, cols=2)
        table.rows[0].cells[0].text = "Databases"
        table.rows[0].cells[1].text = "PostgreSQL, Redis"

        docx_bytes = io.BytesIO()
        doc.save(docx_bytes)
        docx_bytes.seek(0)

        extracted = extract_text_from_docx(docx_bytes)
        self.assertIn("Jane Smith Resume", extracted)
        self.assertIn("TypeScript and Node.js", extracted)
        self.assertIn("PostgreSQL", extracted)

    def test_pdf_parser(self):
        # Create minimal in-memory PDF using pypdf
        writer = pypdf.PdfWriter()
        writer.add_blank_page(width=200, height=200)

        pdf_bytes = io.BytesIO()
        writer.write(pdf_bytes)
        pdf_bytes.seek(0)

        # Blank PDF has no text and should raise DocumentParsingError
        with self.assertRaises(Exception):
            extract_text_from_pdf(pdf_bytes)


if __name__ == "__main__":
    unittest.main()
